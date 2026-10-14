"""Builds VeriBook's document test set: realistic invoice layouts, each in the forms real invoices arrive in.

Every layout is written once as HTML, printed to a born-digital PDF by a headless browser, and from that PDF
turned into the other forms: an image-only scan (clean and messy), a phone photo, a screenshot, and for some
layouts a Word or Excel file. Each layout gets an answer key (truth.json) computed from the same numbers that
printed it, so the key cannot drift from the document.

The forms a reader must cope with, and what each one tests:
  digital.pdf      text layer present: layout and table reading, no OCR
  scan_clean.pdf   image only, 200 dpi, light noise: OCR on a good office scan
  scan_messy.pdf   image only, 150 dpi, skewed, blurred, speckled, JPEG-compressed: OCR on a bad scan
  photo.jpg        a phone photo: perspective, uneven light, a background around the page
  screenshot.png   a clean image of the page
  document.docx    the same invoice as a Word document (some layouts)
  workbook.xlsx    the same invoice as a spreadsheet (some layouts)

Run with the document reader's virtualenv (it has Pillow, numpy, pypdfium2, python-docx, openpyxl):
  docreader/.venv/Scripts/python eval/documents/generate.py
Needs Microsoft Edge or Google Chrome to print the HTML (pass --browser PATH otherwise).

The vendors, GSTINs and VAT numbers are invented; any resemblance to a real business is accidental.
"""

from __future__ import annotations

import argparse
import html
import json
import random
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import numpy as np
import pypdfium2 as pdfium
from PIL import Image, ImageEnhance, ImageFilter

D = Decimal
OUT = Path(__file__).resolve().parent / "invoices"
BROWSERS = [
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    "/usr/bin/google-chrome",
    "/usr/bin/chromium",
]


def q2(x) -> Decimal:
    return D(x).quantize(D("0.01"), rounding=ROUND_HALF_UP)


def q0(x) -> Decimal:
    return D(x).quantize(D("1"), rounding=ROUND_HALF_UP)


def inr(x) -> str:
    """Indian digit grouping: 1,04,135.00"""
    neg = x < 0
    s = f"{abs(q2(x)):.2f}"
    whole, frac = s.split(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        if head:
            groups.insert(0, head)
        whole = ",".join(groups + [tail])
    return ("-" if neg else "") + whole + "." + frac


def west(x) -> str:
    return f"{q2(x):,.2f}"


def e(s) -> str:
    return html.escape(str(s))


@dataclass
class Layout:
    """One invoice layout: its HTML, its answer key, and a plain table view for the Word/Excel forms."""

    name: str
    title: str
    html: str
    truth: dict
    header_lines: list[str] = field(default_factory=list)
    table_header: list[str] = field(default_factory=list)
    table_rows: list[list[str]] = field(default_factory=list)
    totals: list[tuple[str, str]] = field(default_factory=list)
    footer_lines: list[str] = field(default_factory=list)
    forms: tuple[str, ...] = ("digital.pdf", "scan_clean.pdf", "scan_messy.pdf", "photo.jpg")


def truth(vendor, number, date, po, currency, lines, subtotal, tax, total, *, charges=0, discount=0, round_off=0, vendor_tax_id=None, notes=None):
    return {
        "vendor_name": vendor,
        "vendor_tax_id": vendor_tax_id,
        "invoice_number": number,
        "invoice_date": date,
        "po_reference": po,
        "currency": currency,
        "line_items": [{"description": d, "quantity": float(q), "unit_price": float(p)} for d, q, p in lines],
        "subtotal": float(subtotal),
        "tax": float(tax),
        "charges": float(charges),
        "discount": float(discount),
        "round_off": float(round_off),
        "total": float(total),
        "notes": notes,
    }


BASE_CSS = """
@page { size: A4; margin: 14mm; }
* { box-sizing: border-box; }
body { font-family: Arial, Helvetica, sans-serif; font-size: 11px; color: #111; margin: 0; }
table { border-collapse: collapse; width: 100%; }
td, th { vertical-align: top; }
.num { text-align: right; white-space: nowrap; }
.c { text-align: center; }
.b { font-weight: bold; }
.muted { color: #555; }
"""


# ---------------------------------------------------------------------------------------------------------------
# 1. Tally-style GST tax invoice: boxed grid, Buyer's Order No., CGST + SGST rows inside the item table, Round Off,
#    amount in words, and an HSN summary table that must NOT be read as line items.
# ---------------------------------------------------------------------------------------------------------------
def tally_gst() -> Layout:
    vendor = "Shree Ganesh Industrial Supplies"
    gstin = "27AAKFS4821M1Z3"
    items = [
        ("Hex Bolt M10 x 50 (Zinc Plated)", "73181500", D(500), "Nos", D("6.40")),
        ("Spring Washer M10", "73182100", D(500), "Nos", D("0.85")),
        ("Industrial Grease 1 kg (Lithium Complex)", "34031900", D(12), "Nos", D("285.00")),
        ("Safety Gloves, Cut Resistant (Level 5)", "61161000", D(24), "Pairs", D("145.50")),
    ]
    amounts = [q2(q * r) for _, _, q, _, r in items]
    sub = sum(amounts)
    cgst = q2(sub * D("0.09"))
    sgst = q2(sub * D("0.09"))
    pre = sub + cgst + sgst
    total = q0(pre)
    ro = total - pre
    rows = "".join(
        f"<tr><td class='c'>{i}</td><td><b>{e(d)}</b></td><td class='c'>{h}</td><td class='num'><b>{q:,} {u}</b></td>"
        f"<td class='num'>{inr(r)}</td><td class='c'>{u}</td><td class='num'><b>{inr(a)}</b></td></tr>"
        for i, ((d, h, q, u, r), a) in enumerate(zip(items, amounts), 1)
    )
    hsn_rows = "".join(
        f"<tr><td>{h}</td><td class='num'>{inr(a)}</td><td class='c'>9%</td><td class='num'>{inr(q2(a * D('0.09')))}</td>"
        f"<td class='c'>9%</td><td class='num'>{inr(q2(a * D('0.09')))}</td><td class='num'>{inr(q2(a * D('0.18')))}</td></tr>"
        for (_, h, _, _, _), a in zip(items, amounts)
    )
    body = f"""
<style>{BASE_CSS}
.grid td {{ border: 1px solid #000; padding: 3px 5px; }}
.items th {{ border: 1px solid #000; padding: 4px; font-weight: normal; }}
.items td {{ border-left: 1px solid #000; border-right: 1px solid #000; padding: 3px 5px; }}
.items tr.last td {{ border-bottom: 1px solid #000; }}
.items tr.tot td {{ border: 1px solid #000; }}
</style>
<div class='c b' style='font-size:14px'>Tax Invoice</div>
<div style='text-align:right;font-style:italic;font-size:10px'>(ORIGINAL FOR RECIPIENT)</div>
<table class='grid'>
<tr><td rowspan='3' style='width:50%'><b style='font-size:13px'>{vendor}</b><br>Plot 14, MIDC Bhosari, Pune - 411026<br>
GSTIN/UIN: {gstin}<br>State Name : Maharashtra, Code : 27<br>E-Mail : accounts@sgis-pune.example</td>
<td>Invoice No.<br><b>SGIS/2026-27/0412</b></td><td>Dated<br><b>12-Sep-2026</b></td></tr>
<tr><td>Delivery Note<br>DN-0388</td><td>Mode/Terms of Payment<br>30 Days</td></tr>
<tr><td>Buyer's Order No.<br><b>PO-2026-0871</b></td><td>Dated<br>02-Sep-2026</td></tr>
<tr><td rowspan='2'>Buyer (Bill to)<br><b>Northwind Manufacturing Pvt Ltd</b><br>Gat No. 112, Chakan Industrial Area, Pune 410501<br>
GSTIN/UIN : 27AACCN7731Q1Z8<br>State Name : Maharashtra, Code : 27</td>
<td>Dispatched through<br>Own Vehicle</td><td>Destination<br>Chakan</td></tr>
<tr><td colspan='2'>Terms of Delivery<br>Door delivery</td></tr>
</table>
<table class='items'>
<tr><th style='width:5%'>Sl<br>No.</th><th>Description of Goods</th><th>HSN/SAC</th><th>Quantity</th><th>Rate</th><th>per</th><th>Amount</th></tr>
{rows}
<tr><td></td><td style='text-align:right'><b><i>CGST @ 9%</i></b></td><td></td><td></td><td></td><td></td><td class='num'><b>{inr(cgst)}</b></td></tr>
<tr><td></td><td style='text-align:right'><b><i>SGST @ 9%</i></b></td><td></td><td></td><td></td><td></td><td class='num'><b>{inr(sgst)}</b></td></tr>
<tr class='last'><td></td><td style='text-align:right'><i>Round Off</i></td><td></td><td></td><td></td><td></td><td class='num'>{inr(ro)}</td></tr>
<tr class='tot'><td></td><td style='text-align:right'>Total</td><td></td><td class='num'><b>1,036 Nos/Pairs</b></td><td></td><td></td><td class='num b'>&#8377; {inr(total)}</td></tr>
</table>
<table class='grid'><tr><td>Amount Chargeable (in words)<br><b>INR Twelve Thousand Four Hundred Thirty Four Only</b></td><td style='text-align:right'>E. &amp; O.E</td></tr></table>
<table class='grid' style='font-size:10px'>
<tr><td rowspan='2'>HSN/SAC</td><td rowspan='2' class='c'>Taxable<br>Value</td><td colspan='2' class='c'>Central Tax</td><td colspan='2' class='c'>State Tax</td><td rowspan='2' class='c'>Total<br>Tax Amount</td></tr>
<tr><td>Rate</td><td>Amount</td><td>Rate</td><td>Amount</td></tr>
{hsn_rows}
<tr class='b'><td style='text-align:right'>Total</td><td class='num'>{inr(sub)}</td><td></td><td class='num'>{inr(cgst)}</td><td></td><td class='num'>{inr(sgst)}</td><td class='num'>{inr(cgst + sgst)}</td></tr>
</table>
<table class='grid'><tr><td style='width:60%'>Tax Amount (in words) : <b>INR One Thousand Eight Hundred Ninety Six and Sixty Six paise Only</b><br><br>
<u>Declaration</u><br>We declare that this invoice shows the actual price of the goods described and that all particulars are true and correct.</td>
<td style='text-align:right'><b>for {vendor}</b><br><br><br>Authorised Signatory</td></tr></table>
<div class='c muted' style='font-size:9px;margin-top:4px'>This is a Computer Generated Invoice</div>
"""
    t = truth(vendor, "SGIS/2026-27/0412", "2026-09-12", "PO-2026-0871", "INR",
              [(d, q, r) for d, _, q, _, r in items], sub, cgst + sgst, total, round_off=ro, vendor_tax_id=gstin)
    return Layout("tally_gst", "Tally-style GST invoice (CGST+SGST, round off, HSN summary)", body, t,
                  forms=("digital.pdf", "scan_clean.pdf", "scan_messy.pdf", "photo.jpg", "screenshot.png"))


# ---------------------------------------------------------------------------------------------------------------
# 2. Zoho-style inter-state invoice: IGST per line, an "IGST Amt" column next to the line "Amount" (the taxable
#    value), DD/MM/YYYY dates, "P.O.#".
# ---------------------------------------------------------------------------------------------------------------
def zoho_igst() -> Layout:
    vendor = "Kaveri Cloud Solutions LLP"
    gstin = "29AAQFK5520H1ZK"
    items = [
        ("Annual Software Licence - Inventory Module (20 users)", "997331", D(20), D("1850.00")),
        ("Onboarding & Training (remote sessions)", "998311", D(6), D("2500.00")),
        ("Priority Support - Quarterly", "997331", D(1), D("9000.00")),
    ]
    amounts = [q2(q * r) for _, _, q, r in items]
    igst = [q2(a * D("0.18")) for a in amounts]
    sub = sum(amounts)
    tax = sum(igst)
    total = sub + tax
    rows = "".join(
        f"<tr><td class='c'>{i}</td><td>{e(d)}</td><td class='c'>{h}</td><td class='num'>{q:.2f}</td><td class='num'>{inr(r)}</td>"
        f"<td class='c'>18%</td><td class='num'>{inr(g)}</td><td class='num'>{inr(a)}</td></tr>"
        for i, ((d, h, q, r), a, g) in enumerate(zip(items, amounts, igst), 1)
    )
    body = f"""
<style>{BASE_CSS}
body {{ font-family: 'Segoe UI', Arial, sans-serif; }}
.items th {{ background: #3c3d3a; color: #fff; padding: 7px 6px; font-weight: 600; text-align: left; }}
.items td {{ padding: 8px 6px; border-bottom: 1px solid #ddd; }}
.tot td {{ padding: 5px 6px; }}
</style>
<table><tr><td><div style='font-size:20px;font-weight:600;color:#1a5fb4'>Kaveri Cloud</div>
<b>{vendor}</b><br>No. 42, 3rd Floor, 80 Feet Road, Indiranagar<br>Bengaluru Karnataka 560038<br>India<br>GSTIN {gstin}</td>
<td style='text-align:right'><div style='font-size:26px;letter-spacing:1px'>TAX INVOICE</div>
<div class='muted'># KCS-INV-00982</div><br><div class='muted'>Balance Due</div><div class='b' style='font-size:15px'>&#8377;{inr(total)}</div></td></tr></table>
<br>
<table><tr><td style='width:55%'><span class='muted'>Bill To</span><br><b>Northwind Manufacturing Pvt Ltd</b><br>Gat No. 112, Chakan Industrial Area<br>Pune 410501 Maharashtra<br>GSTIN 27AACCN7731Q1Z8<br><br><span class='muted'>Place Of Supply: Maharashtra (27)</span></td>
<td><table>
<tr><td style='text-align:right' class='muted'>Invoice Date :</td><td style='text-align:right'>05/10/2026</td></tr>
<tr><td style='text-align:right' class='muted'>Terms :</td><td style='text-align:right'>Net 15</td></tr>
<tr><td style='text-align:right' class='muted'>Due Date :</td><td style='text-align:right'>20/10/2026</td></tr>
<tr><td style='text-align:right' class='muted'>P.O.# :</td><td style='text-align:right'>PO-2026-0904</td></tr>
</table></td></tr></table>
<br>
<table class='items'><tr><th>#</th><th>Item &amp; Description</th><th>HSN/SAC</th><th class='num'>Qty</th><th class='num'>Rate</th><th>IGST %</th><th class='num'>IGST Amt</th><th class='num'>Amount</th></tr>
{rows}</table>
<table class='tot' style='width:50%;margin-left:50%;margin-top:8px'>
<tr><td style='text-align:right'>Sub Total</td><td class='num'>{inr(sub)}</td></tr>
<tr><td style='text-align:right'>IGST18 (18%)</td><td class='num'>{inr(tax)}</td></tr>
<tr class='b'><td style='text-align:right'>Total</td><td class='num'>&#8377;{inr(total)}</td></tr>
<tr class='b' style='background:#f2f2f2'><td style='text-align:right'>Balance Due</td><td class='num'>&#8377;{inr(total)}</td></tr>
</table>
<p>Total In Words: <i><b>Indian Rupee Seventy-One Thousand Nine Hundred Eighty Only</b></i></p>
<p class='muted'>Notes<br>Thanks for your business. Bank: HDFC Bank, A/c 50200087761234, IFSC HDFC0001234.</p>
"""
    t = truth(vendor, "KCS-INV-00982", "2026-10-05", "PO-2026-0904", "INR", [(d, q, r) for d, _, q, r in items], sub, tax, total, vendor_tax_id=gstin)
    lay = Layout("zoho_igst", "Zoho-style IGST invoice (per-line tax columns, DD/MM dates)", body, t,
                 forms=("digital.pdf", "scan_clean.pdf", "scan_messy.pdf", "photo.jpg", "workbook.xlsx"))
    lay.header_lines = [vendor, "Bengaluru Karnataka 560038", f"GSTIN {gstin}", "TAX INVOICE", "Invoice#: KCS-INV-00982",
                        "Invoice Date: 05/10/2026", "Due Date: 20/10/2026", "P.O.#: PO-2026-0904", "Bill To: Northwind Manufacturing Pvt Ltd"]
    lay.table_header = ["#", "Item & Description", "HSN/SAC", "Qty", "Rate", "IGST %", "IGST Amt", "Amount"]
    lay.table_rows = [[str(i), d, h, f"{q:.2f}", f"{r:.2f}", "18%", f"{g:.2f}", f"{a:.2f}"] for i, ((d, h, q, r), a, g) in enumerate(zip(items, amounts, igst), 1)]
    lay.totals = [("Sub Total", f"{sub:.2f}"), ("IGST18 (18%)", f"{tax:.2f}"), ("Total", f"{total:.2f}")]
    return lay


# ---------------------------------------------------------------------------------------------------------------
# 3. Professional-services bill: no quantity column (amount-only rows), Indian lakh grouping, no PO.
# ---------------------------------------------------------------------------------------------------------------
def service_bill() -> Layout:
    vendor = "Arjun Mehta & Associates"
    gstin = "07AAJFA3391P1ZQ"
    items = [("Professional fees for statutory audit FY 2025-26", D("85000.00")), ("Out of pocket expenses (travel and stationery)", D("3250.00"))]
    sub = sum(a for _, a in items)
    tax = q2(sub * D("0.18"))
    total = sub + tax
    rows = "".join(f"<tr><td class='c'>{i}</td><td>{e(d)}</td><td class='c'>998221</td><td class='num'>{inr(a)}</td></tr>" for i, (d, a) in enumerate(items, 1))
    body = f"""
<style>{BASE_CSS}
body {{ font-family: Georgia, 'Times New Roman', serif; font-size: 12px; }}
.items th, .items td {{ border: 1px solid #444; padding: 6px; }}
</style>
<div class='c'><div style='font-size:20px;font-weight:bold;letter-spacing:1px'>{vendor.upper()}</div>
<div>Chartered Accountants</div><div class='muted'>B-17, Lajpat Nagar II, New Delhi 110024 &nbsp;|&nbsp; GSTIN: {gstin} &nbsp;|&nbsp; PAN: AAJFA3391P</div></div>
<hr>
<div class='c b' style='font-size:15px;margin:6px 0'>INVOICE</div>
<table><tr><td>To,<br><b>Northwind Manufacturing Pvt Ltd</b><br>Gat No. 112, Chakan Industrial Area<br>Pune 410501</td>
<td style='text-align:right'>Bill No: <b>AMA/112/26-27</b><br>Date: 30 September 2026</td></tr></table>
<br>
<table class='items'><tr><th style='width:7%'>S.No.</th><th>Particulars</th><th style='width:12%'>SAC</th><th style='width:20%'>Amount (&#8377;)</th></tr>
{rows}
<tr><td></td><td style='text-align:right'>Taxable Value</td><td></td><td class='num'>{inr(sub)}</td></tr>
<tr><td></td><td style='text-align:right'>IGST @ 18%</td><td></td><td class='num'>{inr(tax)}</td></tr>
<tr class='b'><td></td><td style='text-align:right'>Total Amount Payable</td><td></td><td class='num'>{inr(total)}</td></tr>
</table>
<p>Rupees One Lakh Four Thousand One Hundred Thirty Five Only</p>
<p>Kindly remit to: ICICI Bank, Lajpat Nagar branch, A/c No. 004105012345, IFSC ICIC0000041.<br>Payment due within 15 days of the invoice date.</p>
<p style='text-align:right;margin-top:40px'>For {vendor}<br><br><br>(Partner)<br>M. No. 512873</p>
"""
    t = truth(vendor, "AMA/112/26-27", "2026-09-30", None, "INR", [(d, 1, a) for d, a in items], sub, tax, total, vendor_tax_id=gstin)
    lay = Layout("service_bill", "Professional-services bill (amount-only rows, lakh grouping, no PO)", body, t,
                 forms=("digital.pdf", "scan_clean.pdf", "scan_messy.pdf", "photo.jpg", "document.docx"))
    lay.header_lines = [vendor.upper(), "Chartered Accountants", f"B-17, Lajpat Nagar II, New Delhi 110024 | GSTIN: {gstin}", "INVOICE",
                        "To: Northwind Manufacturing Pvt Ltd, Pune 410501", "Bill No: AMA/112/26-27", "Date: 30 September 2026"]
    lay.table_header = ["S.No.", "Particulars", "SAC", "Amount (₹)"]
    lay.table_rows = [[str(i), d, "998221", inr(a)] for i, (d, a) in enumerate(items, 1)]
    lay.totals = [("Taxable Value", inr(sub)), ("IGST @ 18%", inr(tax)), ("Total Amount Payable", inr(total))]
    lay.footer_lines = ["Rupees One Lakh Four Thousand One Hundred Thirty Five Only", "Payment due within 15 days of the invoice date."]
    return lay


# ---------------------------------------------------------------------------------------------------------------
# 4. US invoice: quantity in the FIRST column, MM/DD/YYYY, sales tax plus a shipping charge.
# ---------------------------------------------------------------------------------------------------------------
def us_invoice() -> Layout:
    vendor = "Bayside Industrial Supply Co."
    items = [
        (D(4), "Pneumatic Cylinder, 50 mm bore x 200 mm stroke", D("189.00")),
        (D(10), 'Air Hose 3/8" x 25 ft, polyurethane', D("32.50")),
        (D(2), "Pressure Regulator with Gauge, 1/4 NPT", D("74.95")),
    ]
    amounts = [q2(q * p) for q, _, p in items]
    sub = sum(amounts)
    tax = q2(sub * D("0.0825"))
    ship = D("45.00")
    total = sub + tax + ship
    rows = "".join(f"<tr><td class='c'>{q}</td><td>{e(d)}</td><td class='num'>{west(p)}</td><td class='num'>{west(a)}</td></tr>" for (q, d, p), a in zip(items, amounts))
    body = f"""
<style>{BASE_CSS}
body {{ font-size: 12px; }}
.items th {{ border-top: 2px solid #234; border-bottom: 2px solid #234; padding: 6px; text-align: left; }}
.items td {{ padding: 6px; border-bottom: 1px solid #ccd; }}
</style>
<table><tr><td><div style='font-size:22px;font-weight:bold;color:#234'>{vendor}</div>
1850 Zanker Road, Suite 140<br>San Jose, CA 95112<br>(408) 555-0182 &middot; billing@baysideindustrial.example</td>
<td style='text-align:right;vertical-align:top'><div style='font-size:30px;color:#234;letter-spacing:2px'>INVOICE</div></td></tr></table>
<br>
<table><tr><td style='width:33%'><b>BILL TO</b><br>Northwind Manufacturing Inc.<br>Accounts Payable<br>2200 Mission College Blvd<br>Santa Clara, CA 95054</td>
<td style='width:33%'><b>SHIP TO</b><br>Northwind Manufacturing Inc.<br>Dock 4<br>2200 Mission College Blvd<br>Santa Clara, CA 95054</td>
<td><table>
<tr><td class='b'>INVOICE #</td><td class='num'>10442</td></tr>
<tr><td class='b'>DATE</td><td class='num'>10/02/2026</td></tr>
<tr><td class='b'>P.O. NUMBER</td><td class='num'>PO-7731</td></tr>
<tr><td class='b'>DUE DATE</td><td class='num'>11/01/2026</td></tr>
</table></td></tr></table>
<br>
<table class='items'><tr><th style='width:9%'>QTY</th><th>DESCRIPTION</th><th class='num'>UNIT PRICE</th><th class='num'>AMOUNT</th></tr>
{rows}</table>
<table style='width:45%;margin-left:55%;margin-top:6px'>
<tr><td>Subtotal</td><td class='num'>{west(sub)}</td></tr>
<tr><td>Sales Tax 8.25%</td><td class='num'>{west(tax)}</td></tr>
<tr><td>Shipping &amp; Handling</td><td class='num'>{west(ship)}</td></tr>
<tr class='b' style='font-size:14px'><td style='border-top:2px solid #234'>TOTAL</td><td class='num' style='border-top:2px solid #234'>${west(total)}</td></tr>
</table>
<p style='margin-top:30px'><b>Terms &amp; Conditions</b><br>Payment is due within 30 days. Please make checks payable to {vendor}.</p>
"""
    t = truth(vendor, "10442", "2026-10-02", "PO-7731", "USD", [(d, q, p) for q, d, p in items], sub, tax, total, charges=ship)
    lay = Layout("us_invoice", "US invoice (qty first, MM/DD dates, sales tax + shipping)", body, t,
                 forms=("digital.pdf", "scan_clean.pdf", "scan_messy.pdf", "photo.jpg", "screenshot.png", "workbook.xlsx"))
    lay.header_lines = [vendor, "1850 Zanker Road, Suite 140, San Jose, CA 95112", "INVOICE", "Invoice #: 10442", "Date: 10/02/2026",
                        "P.O. Number: PO-7731", "Bill To: Northwind Manufacturing Inc."]
    lay.table_header = ["QTY", "DESCRIPTION", "UNIT PRICE", "AMOUNT"]
    lay.table_rows = [[str(q), d, f"{p:.2f}", f"{a:.2f}"] for (q, d, p), a in zip(items, amounts)]
    lay.totals = [("Subtotal", f"{sub:.2f}"), ("Sales Tax 8.25%", f"{tax:.2f}"), ("Shipping & Handling", f"{ship:.2f}"), ("TOTAL", f"{total:.2f}")]
    return lay


# ---------------------------------------------------------------------------------------------------------------
# 5. UK VAT invoice: product codes, Net / VAT / Gross wording, "Your Order Ref".
# ---------------------------------------------------------------------------------------------------------------
def uk_vat() -> Layout:
    vendor = "Pennine Lab Supplies Ltd"
    vat_no = "GB 284 5512 09"
    items = [
        ("LB-220", "Borosilicate Beaker 500 ml (pack of 6)", D(5), D("18.40")),
        ("LB-318", "Nitrile Gloves, Medium (box of 100)", D(12), D("6.75")),
        ("LB-904", "Digital Thermometer -50 to 300 C", D(3), D("24.99")),
    ]
    nets = [q2(q * p) for _, _, q, p in items]
    sub = sum(nets)
    vat = q2(sub * D("0.20"))
    total = sub + vat
    rows = "".join(f"<tr><td>{c}</td><td>{e(d)}</td><td class='num'>{q}</td><td class='num'>&pound;{west(p)}</td><td class='num'>&pound;{west(n)}</td><td class='c'>20%</td></tr>"
                   for (c, d, q, p), n in zip(items, nets))
    body = f"""
<style>{BASE_CSS}
body {{ font-family: Calibri, Arial, sans-serif; font-size: 12px; }}
.items th {{ background: #e6eef5; padding: 6px; text-align: left; border-bottom: 1px solid #7a9cb8; }}
.items td {{ padding: 6px; border-bottom: 1px solid #e3e3e3; }}
</style>
<table><tr><td><div style='font-size:22px;color:#2d5a7b;font-weight:bold'>{vendor}</div>
Unit 7, Calder Park, Wakefield WF2 7AW<br>Tel 01924 555 610 &middot; sales@penninelab.example<br>VAT Reg No: {vat_no} &middot; Company No. 08812274</td>
<td style='text-align:right'><div style='font-size:24px'>Invoice</div></td></tr></table>
<br>
<table><tr><td style='width:55%'><b>Invoice to:</b><br>Northwind Research Ltd<br>Goods In, Building 3<br>Science Park, Cambridge CB4 0WS</td>
<td><table>
<tr><td>Invoice Number:</td><td class='num b'>PLS-55821</td></tr>
<tr><td>Invoice Date:</td><td class='num'>14 Oct 2026</td></tr>
<tr><td>Your Order Ref:</td><td class='num'>PO-UK-2219</td></tr>
<tr><td>Account No:</td><td class='num'>NOR001</td></tr>
</table></td></tr></table>
<br>
<table class='items'><tr><th>Code</th><th>Description</th><th class='num'>Qty</th><th class='num'>Unit Price</th><th class='num'>Net</th><th class='c'>VAT %</th></tr>
{rows}</table>
<table style='width:40%;margin-left:60%;margin-top:8px'>
<tr><td>Net Total</td><td class='num'>&pound;{west(sub)}</td></tr>
<tr><td>VAT @ 20%</td><td class='num'>&pound;{west(vat)}</td></tr>
<tr class='b'><td>Invoice Total</td><td class='num'>&pound;{west(total)}</td></tr>
</table>
<p class='muted' style='margin-top:30px'>Payment terms: 30 days from invoice date. Bank: Barclays, Sort Code 20-45-77, Account 43918822.</p>
"""
    t = truth(vendor, "PLS-55821", "2026-10-14", "PO-UK-2219", "GBP", [(d, q, p) for _, d, q, p in items], sub, vat, total, vendor_tax_id=vat_no)
    lay = Layout("uk_vat", "UK VAT invoice (codes, Net/VAT/Gross, Your Order Ref)", body, t,
                 forms=("digital.pdf", "scan_clean.pdf", "scan_messy.pdf", "photo.jpg", "document.docx"))
    lay.header_lines = [vendor, "Unit 7, Calder Park, Wakefield WF2 7AW", f"VAT Reg No: {vat_no}", "Invoice", "Invoice Number: PLS-55821",
                        "Invoice Date: 14 Oct 2026", "Your Order Ref: PO-UK-2219", "Invoice to: Northwind Research Ltd"]
    lay.table_header = ["Code", "Description", "Qty", "Unit Price", "Net", "VAT %"]
    lay.table_rows = [[c, d, str(q), f"£{p:.2f}", f"£{n:.2f}", "20%"] for (c, d, q, p), n in zip(items, nets)]
    lay.totals = [("Net Total", f"£{sub:.2f}"), ("VAT @ 20%", f"£{vat:.2f}"), ("Invoice Total", f"£{total:.2f}")]
    return lay


# ---------------------------------------------------------------------------------------------------------------
# 6. A long invoice: 28 lines over two pages, the header row repeated, DD.MM.YYYY date, mixed units.
# ---------------------------------------------------------------------------------------------------------------
def multipage_gst() -> Layout:
    vendor = "Deccan Electricals & Controls"
    gstin = "33AAGCD6612K1Z2"
    catalogue = [
        ("FR PVC Insulated Copper Cable 1.5 sq mm", "Mtr", "18.60"), ("FR PVC Insulated Copper Cable 2.5 sq mm", "Mtr", "29.40"),
        ("FR PVC Insulated Copper Cable 4 sq mm", "Mtr", "46.75"), ("Armoured Cable 4C x 16 sq mm Al", "Mtr", "212.00"),
        ("MCB 6A Single Pole C-Curve", "Nos", "168.00"), ("MCB 16A Single Pole C-Curve", "Nos", "172.00"),
        ("MCB 32A Triple Pole C-Curve", "Nos", "1,145.00"), ("RCCB 40A 4 Pole 30mA", "Nos", "3,480.00"),
        ("Power Contactor 9A 230V AC Coil", "Nos", "1,260.00"), ("Power Contactor 25A 230V AC Coil", "Nos", "2,195.00"),
        ("Thermal Overload Relay 7-10A", "Nos", "1,540.00"), ("Push Button Green (NO)", "Nos", "96.00"),
        ("Push Button Red (NC)", "Nos", "96.00"), ("Pilot Lamp LED 22mm Red", "Nos", "74.00"),
        ("Pilot Lamp LED 22mm Green", "Nos", "74.00"), ("DIN Rail 35mm Slotted (1 m)", "Nos", "118.00"),
        ("Cable Tie 200mm Black (pack of 100)", "Pkt", "86.00"), ("Lugs Copper Ring Type 2.5 sq mm (pack of 100)", "Pkt", "310.00"),
        ("Lugs Copper Ring Type 6 sq mm (pack of 100)", "Pkt", "545.00"), ("PVC Conduit Pipe 25mm (3 m)", "Nos", "92.00"),
        ("Junction Box 4 Way 25mm", "Nos", "38.00"), ("Distribution Board 8 Way TPN", "Nos", "4,350.00"),
        ("Rotary Isolator 32A 4 Pole", "Nos", "1,980.00"), ("Terminal Block 4 sq mm Grey", "Nos", "28.50"),
        ("Terminal Block End Plate", "Nos", "9.00"), ("Ferrule Marker Set 0-9 (1.5 sq mm)", "Box", "245.00"),
        ("Insulation Tape 18mm x 8m Black", "Nos", "22.00"), ("Panel Cooling Fan 120mm 230V", "Nos", "865.00"),
    ]
    qtys = [D(300), D(250), D(150), D(80), D(24), D(24), D(6), D(4), D(10), D(6), D(10), D(30), D(30), D(40), D(40), D(20), D(15), D(5), D(5),
            D(60), D(45), D(3), D(4), D(120), D(30), D(6), D(50), D(8)]
    items = [(d, u, D(r.replace(",", "")), q) for (d, u, r), q in zip(catalogue, qtys)]
    amounts = [q2(q * r) for _, _, r, q in items]
    sub = sum(amounts)
    cgst = q2(sub * D("0.09"))
    sgst = q2(sub * D("0.09"))
    pre = sub + cgst + sgst
    total = q0(pre)
    ro = total - pre
    rows = "".join(f"<tr><td class='c'>{i}</td><td>{e(d)}</td><td class='c'>8544</td><td class='num'>{q}</td><td class='c'>{u}</td><td class='num'>{inr(r)}</td><td class='num'>{inr(a)}</td></tr>"
                   for i, ((d, u, r, q), a) in enumerate(zip(items, amounts), 1))
    body = f"""
<style>{BASE_CSS}
body {{ font-size: 12.5px; }}
.items th {{ border: 1px solid #333; background: #eee; padding: 9px 5px; }}
.items td {{ border: 1px solid #999; padding: 10px 5px; }}
thead {{ display: table-header-group; }}
tr {{ page-break-inside: avoid; }}
</style>
<table><tr><td><div style='font-size:18px;font-weight:bold'>{vendor}</div>No. 8, Sidco Industrial Estate, Guindy, Chennai 600032<br>GSTIN: {gstin} &middot; Ph: 044-2250 1187</td>
<td style='text-align:right'><div style='font-size:18px;font-weight:bold'>TAX INVOICE</div>Invoice No.: <b>DEC/26/1187</b><br>Invoice Date: 08.10.2026<br>Purchase Order No.: PO-2026-0950</td></tr></table>
<p>Consignee: <b>Northwind Manufacturing Pvt Ltd</b>, Plant 2, Sriperumbudur, Tamil Nadu 602105 (GSTIN 33AACCN7731Q1Z1)</p>
<table class='items'><thead><tr><th>Sr</th><th>Description</th><th>HSN</th><th>Qty</th><th>Unit</th><th>Rate</th><th>Amount</th></tr></thead>
<tbody>{rows}</tbody></table>
<table style='width:45%;margin-left:55%;margin-top:8px'>
<tr><td>Total Taxable Value</td><td class='num'>{inr(sub)}</td></tr>
<tr><td>CGST 9%</td><td class='num'>{inr(cgst)}</td></tr>
<tr><td>SGST 9%</td><td class='num'>{inr(sgst)}</td></tr>
<tr><td>Rounding</td><td class='num'>{inr(ro)}</td></tr>
<tr class='b'><td>Grand Total</td><td class='num'>&#8377; {inr(total)}</td></tr>
</table>
<p>Goods once sold will not be taken back. Subject to Chennai jurisdiction.</p>
<p style='text-align:right'>For {vendor}<br><br>Authorised Signatory</p>
"""
    t = truth(vendor, "DEC/26/1187", "2026-10-08", "PO-2026-0950", "INR", [(d, q, r) for d, _, r, q in items], sub, cgst + sgst, total, round_off=ro, vendor_tax_id=gstin)
    return Layout("multipage_gst", "Two-page GST invoice (28 lines, repeated header, DD.MM.YYYY)", body, t,
                  forms=("digital.pdf", "scan_clean.pdf", "scan_messy.pdf"))


# ---------------------------------------------------------------------------------------------------------------
# 7. Trader's bill: line discounts, decimal quantities in KG, a per-line GST column and a per-line TOTAL column
#    (tax included) next to the taxable value, negative round off.
# ---------------------------------------------------------------------------------------------------------------
def discount_uom() -> Layout:
    vendor = "Malabar Spice Traders"
    gstin = "32ABMPT7740C1ZD"
    items = [
        ("Black Pepper Whole (Grade A, 550 GL)", "09041110", D("25.000"), "KG", D("640.00"), D("5")),
        ("Green Cardamom 8 mm Bold", "09083120", D("4.500"), "KG", D("2450.00"), D("0")),
        ("Turmeric Powder 1 kg Pouch", "09103030", D("40"), "PCS", D("118.00"), D("2.5")),
    ]
    rows_data = []
    for d, h, q, u, r, dp in items:
        gross = q2(q * r)
        disc = q2(gross * dp / 100)
        taxable = gross - disc
        gst = q2(taxable * D("0.05"))
        rows_data.append((d, h, q, u, r, dp, taxable, gst, taxable + gst))
    sub = sum(x[6] for x in rows_data)
    gst_total = sum(x[7] for x in rows_data)
    cgst = q2(gst_total / 2)
    sgst = gst_total - cgst
    pre = sub + gst_total
    total = q0(pre)
    ro = total - pre
    rows = "".join(
        f"<tr><td class='c'>{i}</td><td>{e(d)}</td><td class='c'>{h}</td><td class='num'>{q}</td><td class='c'>{u}</td><td class='num'>{inr(r)}</td>"
        f"<td class='c'>{dp}%</td><td class='num'>{inr(tv)}</td><td class='c'>5%</td><td class='num'>{inr(g)}</td><td class='num'>{inr(lt)}</td></tr>"
        for i, (d, h, q, u, r, dp, tv, g, lt) in enumerate(rows_data, 1)
    )
    body = f"""
<style>{BASE_CSS}
body {{ font-size: 10.5px; }}
.items th {{ border: 1px solid #555; padding: 4px; background: #f6f1e7; }}
.items td {{ border: 1px solid #aaa; padding: 4px; }}
</style>
<div class='c'><div style='font-size:20px;font-weight:bold;color:#7a3e00'>{vendor}</div>
Wholesale Dealers in Spices &amp; Condiments<br>Mattancherry, Kochi, Kerala 682002 &middot; GSTIN {gstin} &middot; FSSAI 11319007000412</div>
<div class='c b' style='margin:8px 0;font-size:13px;border-top:1px solid #000;border-bottom:1px solid #000'>GST INVOICE</div>
<table><tr><td>Billed To: <b>Northwind Foods Pvt Ltd</b><br>Edappally, Kochi 682024<br>GSTIN 32AADCN4410F1Z6</td>
<td style='text-align:right'>Invoice No: <b>MST/1029</b><br>Date: 11-10-2026<br>PO No.: PO-2026-0962<br>Transport: KL-07-BX-4471</td></tr></table>
<br>
<table class='items'><tr><th>S.No</th><th>Item</th><th>HSN</th><th>Qty</th><th>UoM</th><th>Rate</th><th>Disc %</th><th>Taxable Value</th><th>GST %</th><th>GST Amt</th><th>Total</th></tr>
{rows}</table>
<table style='width:45%;margin-left:55%;margin-top:8px'>
<tr><td>Taxable Amount</td><td class='num'>{inr(sub)}</td></tr>
<tr><td>CGST @ 2.5%</td><td class='num'>{inr(cgst)}</td></tr>
<tr><td>SGST @ 2.5%</td><td class='num'>{inr(sgst)}</td></tr>
<tr><td>Round Off</td><td class='num'>{inr(ro)}</td></tr>
<tr class='b'><td>Net Payable</td><td class='num'>&#8377; {inr(total)}</td></tr>
</table>
<p>Rupees Thirty Two Thousand Three Hundred Sixty Eight Only</p>
"""
    # VeriBook's line price is what is actually charged per unit: the taxable value / quantity (after the line discount).
    t = truth(vendor, "MST/1029", "2026-10-11", "PO-2026-0962", "INR",
              [(d, q, q2(tv / q)) for d, _, q, _, _, _, tv, _, _ in rows_data], sub, gst_total, total, round_off=ro, vendor_tax_id=gstin,
              notes="unit_price is the net rate after the line discount (taxable value / quantity)")
    return Layout("discount_uom", "Trader's GST bill (line discounts, KG quantities, per-line tax and total columns)", body, t)


# ---------------------------------------------------------------------------------------------------------------
# 8. A small vendor's plain bill: no table at all ("1000 nos @ 1.20 = 1,200.00"), ordinal date, no tax.
# ---------------------------------------------------------------------------------------------------------------
def plain_bill() -> Layout:
    vendor = "Rohan Print Studio"
    items = [("Visiting cards, premium matte 350 gsm", D(1000), D("1.20")), ("Flex banner 6 x 3 ft with eyelets", D(2), D("450.00")),
             ("Letterheads A4, 100 gsm", D(500), D("2.80"))]
    amounts = [q2(q * r) for _, q, r in items]
    total = sum(amounts)
    lines = "".join(f"<p>{i}. {e(d)} - {q} nos @ {inr(r)} = {inr(a)}</p>" for i, ((d, q, r), a) in enumerate(zip(items, amounts), 1))
    body = f"""
<style>{BASE_CSS} body {{ font-family: 'Courier New', monospace; font-size: 13px; }} p {{ margin: 5px 0; }}</style>
<p><b>From: {vendor}</b><br>Shop 3, Ganesh Peth, Nagpur 440018<br>Mob: 98220 41175</p>
<p>To: Northwind Manufacturing Pvt Ltd, Nagpur office</p>
<p>Invoice No. RPS-0093<br>Date: 3rd October 2026</p>
<p>----------------------------------------------</p>
{lines}
<p>----------------------------------------------</p>
<p><b>Total: Rs. {inr(total)}</b></p>
<p>(Not registered under GST. No tax charged.)</p>
<p>Payment by UPI to rohanprint@okaxis or cash.</p>
"""
    t = truth(vendor, "RPS-0093", "2026-10-03", None, "INR", items, total, 0, total)
    lay = Layout("plain_bill", "Small vendor's plain bill (no table, '@ rate = amount' lines, no tax)", body, t,
                 forms=("digital.pdf", "scan_clean.pdf", "scan_messy.pdf", "photo.jpg", "screenshot.png", "document.docx"))
    lay.header_lines = [f"From: {vendor}", "Shop 3, Ganesh Peth, Nagpur 440018", "To: Northwind Manufacturing Pvt Ltd", "Invoice No. RPS-0093", "Date: 3rd October 2026"]
    lay.footer_lines = [f"{i}. {d} - {q} nos @ {inr(r)} = {inr(a)}" for i, ((d, q, r), a) in enumerate(zip(items, amounts), 1)] + [
        f"Total: Rs. {inr(total)}", "(Not registered under GST. No tax charged.)"]
    return lay


LAYOUTS = [tally_gst, zoho_igst, service_bill, us_invoice, uk_vat, multipage_gst, discount_uom, plain_bill]


# ---------------------------------------------------------------------------------------------------------------
# Rendering and degrading
# ---------------------------------------------------------------------------------------------------------------
def find_browser(explicit: str | None) -> str:
    for candidate in ([explicit] if explicit else BROWSERS):
        if candidate and Path(candidate).exists():
            return candidate
    raise SystemExit("No Edge/Chrome found to print the invoices; pass --browser PATH")


def print_pdf(browser: str, html_text: str, out: Path) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        page = Path(tmp) / "invoice.html"
        page.write_text(f"<!doctype html><html><head><meta charset='utf-8'></head><body>{html_text}</body></html>", encoding="utf-8")
        cmd = [browser, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--print-to-pdf-no-header",
               f"--user-data-dir={Path(tmp) / 'profile'}", f"--print-to-pdf={out}", page.as_uri()]
        subprocess.run(cmd, check=True, capture_output=True, timeout=120)
    if not out.exists() or out.stat().st_size == 0:
        raise RuntimeError(f"the browser did not write {out}")


def render_pages(pdf: Path, dpi: int, mode: str = "L") -> list[Image.Image]:
    doc = pdfium.PdfDocument(str(pdf))
    try:
        return [doc[i].render(scale=dpi / 72).to_pil().convert(mode) for i in range(len(doc))]
    finally:
        doc.close()


def noise(img: Image.Image, sigma: float, rng: np.random.Generator) -> Image.Image:
    a = np.asarray(img).astype(np.float32)
    a += rng.normal(0, sigma, a.shape)
    return Image.fromarray(np.clip(a, 0, 255).astype(np.uint8), img.mode)


def specks(img: Image.Image, count: int, rng: np.random.Generator) -> Image.Image:
    a = np.asarray(img).copy()
    h, w = a.shape[:2]
    ys, xs = rng.integers(0, h, count), rng.integers(0, w, count)
    a[ys, xs] = 30
    return Image.fromarray(a, img.mode)


def jpeg_roundtrip(img: Image.Image, quality: int) -> Image.Image:
    import io
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=quality)
    buf.seek(0)
    return Image.open(buf).convert(img.mode)


def save_image_pdf(pages: list[Image.Image], out: Path, dpi: int) -> None:
    pages[0].save(out, "PDF", resolution=dpi, save_all=True, append_images=pages[1:])


def scan_clean(pdf: Path, out: Path, rng) -> None:
    pages = [noise(p, 4, rng) for p in render_pages(pdf, 200)]
    save_image_pdf(pages, out, 200)


def scan_messy(pdf: Path, out: Path, rng) -> None:
    pages = []
    for p in render_pages(pdf, 150):
        p = p.rotate(float(rng.uniform(1.2, 2.2)) * (1 if rng.random() < 0.5 else -1), resample=Image.BICUBIC, expand=True, fillcolor=255)
        p = p.filter(ImageFilter.GaussianBlur(0.6))
        p = ImageEnhance.Contrast(p).enhance(0.8)
        p = ImageEnhance.Brightness(p).enhance(0.93)
        p = specks(noise(p, 11, rng), 1500, rng)
        pages.append(jpeg_roundtrip(p, 40))
    save_image_pdf(pages, out, 150)


def _perspective_coeffs(src, dst):
    """Coefficients for PIL's PERSPECTIVE transform mapping output points dst -> input points src."""
    m = []
    for (x, y), (u, v) in zip(dst, src):
        m.append([x, y, 1, 0, 0, 0, -u * x, -u * y])
        m.append([0, 0, 0, x, y, 1, -v * x, -v * y])
    a = np.array(m, dtype=np.float64)
    b = np.array([c for pt in src for c in pt], dtype=np.float64)
    return np.linalg.solve(a, b).tolist()


def photo(pdf: Path, out: Path, rng) -> None:
    page = render_pages(pdf, 130, "RGB")[0]
    w, h = page.size
    pad = int(w * 0.12)
    canvas = Image.new("RGB", (w + 2 * pad, h + 2 * pad), (92, 74, 58))  # a wooden desk
    canvas = noise(canvas, 9, rng)
    canvas.paste(page, (pad, pad))
    W, H = canvas.size
    # the page photographed slightly from below and to the left: corners pulled in unevenly
    src = [(0, 0), (W, 0), (W, H), (0, H)]
    dst = [(W * 0.035, H * 0.02), (W * 0.975, H * 0.045), (W * 0.99, H * 0.985), (W * 0.01, H * 0.96)]
    canvas = canvas.transform((W, H), Image.PERSPECTIVE, _perspective_coeffs(src, dst), Image.BICUBIC, fillcolor=(70, 56, 44))
    canvas = canvas.rotate(2.5, resample=Image.BICUBIC, fillcolor=(70, 56, 44))
    # uneven light: brighter at the top right, a shadow toward the bottom left
    yy, xx = np.mgrid[0:H, 0:W].astype(np.float32)
    light = 0.72 + 0.33 * ((xx / W) * 0.6 + (1 - yy / H) * 0.4)
    a = np.asarray(canvas).astype(np.float32) * light[..., None]
    canvas = Image.fromarray(np.clip(a, 0, 255).astype(np.uint8), "RGB").filter(ImageFilter.GaussianBlur(0.8))
    canvas.save(out, "JPEG", quality=72)


def screenshot(pdf: Path, out: Path, rng) -> None:
    render_pages(pdf, 110, "RGB")[0].save(out, "PNG")


def write_docx(lay: Layout, out: Path) -> None:
    from docx import Document

    doc = Document()
    for line in lay.header_lines:
        doc.add_paragraph(line)
    if lay.table_header:
        table = doc.add_table(rows=1, cols=len(lay.table_header))
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, lay.table_header):
            cell.text = text
        for row in lay.table_rows:
            for cell, text in zip(table.add_row().cells, row):
                cell.text = text
        for label, value in lay.totals:
            cells = table.add_row().cells
            cells[1].text = label
            cells[-1].text = value
    for line in lay.footer_lines:
        doc.add_paragraph(line)
    doc.save(out)


def write_xlsx(lay: Layout, out: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    for line in lay.header_lines:
        ws.append([line])
    ws.append([])
    ws.append(lay.table_header)
    for row in lay.table_rows:
        ws.append(row)
    for label, value in lay.totals:
        ws.append([""] * (len(lay.table_header) - 2) + [label, value])
    wb.save(out)


FORM_WRITERS = {"scan_clean.pdf": scan_clean, "scan_messy.pdf": scan_messy, "photo.jpg": photo, "screenshot.png": screenshot}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--browser", help="path to Edge or Chrome")
    ap.add_argument("--only", nargs="*", help="layout names to (re)build")
    args = ap.parse_args()
    browser = find_browser(args.browser)
    rng = np.random.default_rng(20261013)  # fixed: the same files every run
    random.seed(20261013)

    index = []
    for make in LAYOUTS:
        lay = make()
        if args.only and lay.name not in args.only:
            continue
        folder = OUT / lay.name
        if folder.exists():
            shutil.rmtree(folder)
        folder.mkdir(parents=True)
        digital = folder / "digital.pdf"
        print_pdf(browser, lay.html, digital)
        for form in lay.forms:
            if form == "digital.pdf":
                continue
            target = folder / form
            if form == "document.docx":
                write_docx(lay, target)
            elif form == "workbook.xlsx":
                write_xlsx(lay, target)
            else:
                FORM_WRITERS[form](digital, target, rng)
        (folder / "truth.json").write_text(json.dumps(lay.truth, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        index.append({"layout": lay.name, "title": lay.title, "forms": list(lay.forms)})
        print(f"{lay.name}: {', '.join(lay.forms)}")
    if not args.only:
        (OUT / "index.json").write_text(json.dumps(index, indent=2) + "\n", encoding="utf-8")
    if not args.only or any(n.startswith("po_") for n in args.only):
        build_purchase_orders(browser, rng)
    if not args.only or any(n.startswith("gr_") for n in args.only):
        build_goods_receipts(browser, rng)


# ---------------------------------------------------------------------------------------------------------------
# Purchase orders: issued BY the buyer (Northwind, the letterhead) TO a vendor (named in a "Vendor" block). The PO
# reader must take the vendor from that block, never the letterhead. Truth: vendor_name, po_number, order_date,
# currency, line_items.
# ---------------------------------------------------------------------------------------------------------------
PO_OUT = Path(__file__).resolve().parent / "purchase_orders"


def po_indian() -> Layout:
    items = [("Hex Bolt M10 x 50 (Zinc Plated)", "Nos", D(500), D("6.40")), ("Spring Washer M10", "Nos", D(500), D("0.85")),
             ("Industrial Grease 1 kg (Lithium Complex)", "Nos", D(12), D("285.00")), ("Safety Gloves, Cut Resistant (Level 5)", "Pairs", D(24), D("145.50"))]
    amounts = [q2(q * r) for _, _, q, r in items]
    sub = sum(amounts)
    rows = "".join(f"<tr><td class='c'>{i}</td><td>{e(d)}</td><td class='num'>{q}</td><td class='c'>{u}</td><td class='num'>{inr(r)}</td><td class='num'>{inr(a)}</td></tr>"
                   for i, ((d, u, q, r), a) in enumerate(zip(items, amounts), 1))
    body = f"""
<style>{BASE_CSS} .items th {{ border: 1px solid #333; background: #eef2f7; padding: 5px; }} .items td {{ border: 1px solid #999; padding: 5px; }}</style>
<table><tr><td><div style='font-size:20px;font-weight:bold;color:#1f3a5f'>Northwind Manufacturing Pvt Ltd</div>
Gat No. 112, Chakan Industrial Area, Pune 410501 &middot; GSTIN 27AACCN7731Q1Z8</td>
<td style='text-align:right'><div style='font-size:20px;font-weight:bold'>PURCHASE ORDER</div>PO Number: <b>PO-2026-0871</b><br>PO Date: 02-Sep-2026</td></tr></table>
<br>
<table><tr><td style='width:50%'><b>Vendor:</b> Shree Ganesh Industrial Supplies<br>Plot 14, MIDC Bhosari, Pune - 411026<br>GSTIN 27AAKFS4821M1Z3</td>
<td><b>Ship To:</b> Northwind Manufacturing Pvt Ltd, Stores, Chakan</td></tr></table>
<br>
<table class='items'><tr><th>Sr</th><th>Description</th><th>Qty</th><th>Unit</th><th>Rate</th><th>Amount</th></tr>{rows}</table>
<table style='width:40%;margin-left:60%;margin-top:6px'><tr><td>Sub Total</td><td class='num'>{inr(sub)}</td></tr>
<tr><td>GST (as applicable)</td><td class='num'>extra</td></tr><tr class='b'><td>PO Value (before tax)</td><td class='num'>&#8377; {inr(sub)}</td></tr></table>
<p>Delivery within 7 days. Please quote the PO number on your invoice.</p>
<p style='text-align:right'>For Northwind Manufacturing Pvt Ltd<br><br>Purchase Manager</p>
"""
    t = {"vendor_name": "Shree Ganesh Industrial Supplies", "po_number": "PO-2026-0871", "order_date": "2026-09-02", "currency": "INR",
         "line_items": [{"description": d, "quantity": float(q), "unit_price": float(r)} for d, _, q, r in items]}
    return Layout("po_indian", "Indian purchase order (vendor block, buyer letterhead)", body, t, forms=("digital.pdf", "scan_clean.pdf", "photo.jpg"))


def po_us() -> Layout:
    items = [(D(4), "Pneumatic Cylinder, 50 mm bore x 200 mm stroke", D("189.00")), (D(10), 'Air Hose 3/8" x 25 ft, polyurethane', D("32.50")),
             (D(2), "Pressure Regulator with Gauge, 1/4 NPT", D("74.95"))]
    amounts = [q2(q * p) for q, _, p in items]
    rows = "".join(f"<tr><td>{e(d)}</td><td class='num'>{q}</td><td class='num'>{west(p)}</td><td class='num'>{west(a)}</td></tr>" for (q, d, p), a in zip(items, amounts))
    body = f"""
<style>{BASE_CSS} body {{ font-size: 12px; }} .items th {{ border-bottom: 2px solid #234; padding: 6px; text-align: left; }} .items td {{ padding: 6px; border-bottom: 1px solid #ccd; }}</style>
<div style='font-size:22px;font-weight:bold;color:#234'>Northwind Manufacturing Inc.</div><div>2200 Mission College Blvd, Santa Clara, CA 95054</div>
<h2 style='margin:14px 0 4px'>Purchase Order PO-7731</h2><div>Order Date: 09/18/2026</div>
<br>
<table><tr><td style='width:50%'><b>Supplier:</b> Bayside Industrial Supply Co.<br>1850 Zanker Road, Suite 140, San Jose, CA 95112</td><td><b>Deliver to:</b> Dock 4, Santa Clara</td></tr></table>
<br>
<table class='items'><tr><th>Description</th><th class='num'>Quantity</th><th class='num'>Unit Price</th><th class='num'>Line Total</th></tr>{rows}</table>
<p style='text-align:right'><b>Order Total: ${west(sum(amounts))}</b> (excl. tax and shipping)</p>
"""
    t = {"vendor_name": "Bayside Industrial Supply Co.", "po_number": "PO-7731", "order_date": "2026-09-18", "currency": "USD",
         "line_items": [{"description": d, "quantity": float(q), "unit_price": float(p)} for q, d, p in items]}
    return Layout("po_us", "US purchase order (Supplier block, MM/DD date)", body, t, forms=("digital.pdf", "scan_messy.pdf"))


# ---------------------------------------------------------------------------------------------------------------
# Goods receipts: the proof that goods arrived. Three real shapes: a vendor's delivery challan (letterhead = vendor,
# quantities dispatched, signed "received in good condition"), a buyer's goods receipt note (ordered / received /
# accepted / rejected columns), and a US packing slip (ordered / shipped / backordered). Truth: po_number,
# document_number, receipt_date, receiver_name, line_items [{description, quantity}] where quantity is what counts
# as received (accepted when there is an accepted column, else shipped/dispatched).
# ---------------------------------------------------------------------------------------------------------------
GR_OUT = Path(__file__).resolve().parent / "goods_receipts"


def gr_challan() -> Layout:
    items = [("Hex Bolt M10 x 50 (Zinc Plated)", "73181500", D(500), "Nos"), ("Spring Washer M10", "73182100", D(500), "Nos"),
             ("Industrial Grease 1 kg (Lithium Complex)", "34031900", D(12), "Nos"), ("Safety Gloves, Cut Resistant (Level 5)", "61161000", D(20), "Pairs")]
    rows = "".join(f"<tr><td class='c'>{i}</td><td>{e(d)}</td><td class='c'>{h}</td><td class='num'>{q}</td><td class='c'>{u}</td></tr>"
                   for i, (d, h, q, u) in enumerate(items, 1))
    body = f"""
<style>{BASE_CSS} .grid td {{ border: 1px solid #000; padding: 4px 6px; }} .items th, .items td {{ border: 1px solid #000; padding: 5px; }}</style>
<div class='c b' style='font-size:15px'>DELIVERY CHALLAN</div>
<table class='grid'><tr><td style='width:55%' rowspan='2'><b style='font-size:13px'>Shree Ganesh Industrial Supplies</b><br>Plot 14, MIDC Bhosari, Pune - 411026<br>GSTIN: 27AAKFS4821M1Z3</td>
<td>Challan No.<br><b>SGIS/DC/0388</b></td><td>Date<br><b>10-Sep-2026</b></td></tr>
<tr><td>Buyer's Order No.<br><b>PO-2026-0871</b></td><td>Vehicle No.<br>MH-14-HJ-2231</td></tr>
<tr><td colspan='3'>Consignee: <b>Northwind Manufacturing Pvt Ltd</b>, Gat No. 112, Chakan Industrial Area, Pune 410501</td></tr></table>
<br>
<table class='items'><tr><th>Sl</th><th>Description of Goods</th><th>HSN</th><th>Quantity</th><th>Unit</th></tr>{rows}</table>
<p>Balance 4 Pairs of Safety Gloves to follow. Goods sent for: Sale against PO.</p>
<table style='margin-top:30px'><tr><td style='width:50%'>Received the above goods in good condition.<br><br>Received by: R. Patil (Stores)<br>Signature &amp; Stamp</td>
<td style='text-align:right'>for Shree Ganesh Industrial Supplies<br><br><br>Authorised Signatory</td></tr></table>
"""
    t = {"po_number": "PO-2026-0871", "document_number": "SGIS/DC/0388", "receipt_date": "2026-09-10", "receiver_name": "R. Patil (Stores)",
         "line_items": [{"description": d, "quantity": float(q)} for d, _, q, _ in items]}
    return Layout("gr_challan", "Vendor's delivery challan (partial: 20 of 24 gloves)", body, t, forms=("digital.pdf", "scan_clean.pdf", "photo.jpg"))


def gr_grn() -> Layout:
    items = [("Black Pepper Whole (Grade A, 550 GL)", "KG", D("25.000"), D("25.000"), D("24.000"), D("1.000"), "1 kg bag torn"),
             ("Green Cardamom 8 mm Bold", "KG", D("4.500"), D("4.500"), D("4.500"), D("0"), ""),
             ("Turmeric Powder 1 kg Pouch", "PCS", D(40), D(40), D(40), D(0), "")]
    rows = "".join(f"<tr><td class='c'>{i}</td><td>{e(d)}</td><td class='c'>{u}</td><td class='num'>{o}</td><td class='num'>{r}</td><td class='num'>{a}</td><td class='num'>{j}</td><td>{e(m)}</td></tr>"
                   for i, (d, u, o, r, a, j, m) in enumerate(items, 1))
    body = f"""
<style>{BASE_CSS} .items th {{ border: 1px solid #333; background: #eef2f7; padding: 5px; }} .items td {{ border: 1px solid #999; padding: 5px; }}</style>
<table><tr><td><div style='font-size:18px;font-weight:bold;color:#1f3a5f'>Northwind Foods Pvt Ltd</div>Edappally, Kochi 682024 &middot; Stores &amp; Receiving</td>
<td style='text-align:right'><div style='font-size:18px;font-weight:bold'>GOODS RECEIPT NOTE</div>GRN No.: <b>NW/GRN/2026/1142</b><br>GRN Date: 11-10-2026</td></tr></table>
<br>
<table><tr><td style='width:50%'>Supplier: <b>Malabar Spice Traders</b><br>Mattancherry, Kochi 682002</td><td>PO No.: <b>PO-2026-0962</b><br>Supplier Invoice: MST/1029</td></tr></table>
<br>
<table class='items'><tr><th>Sr</th><th>Item</th><th>UoM</th><th>Ordered</th><th>Received</th><th>Accepted</th><th>Rejected</th><th>Remarks</th></tr>{rows}</table>
<table style='margin-top:30px;width:100%'><tr><td>Inspected by: S. Nair (QC)</td><td>Received by: R. Patil (Stores)</td><td style='text-align:right'>Store In-charge</td></tr></table>
"""
    t = {"po_number": "PO-2026-0962", "document_number": "NW/GRN/2026/1142", "receipt_date": "2026-10-11", "receiver_name": "R. Patil (Stores)",
         "line_items": [{"description": d, "quantity": float(a)} for d, _, _, _, a, _, _ in items], "rejected": [float(j) for *_, j, _ in items]}
    return Layout("gr_grn", "Buyer's goods receipt note (ordered / received / accepted / rejected)", body, t, forms=("digital.pdf", "scan_messy.pdf"))


def gr_packing_slip() -> Layout:
    items = [("CYL-50-200", "Pneumatic Cylinder, 50 mm bore x 200 mm stroke", 4, 4, 0), ("AH-38-25", 'Air Hose 3/8" x 25 ft, polyurethane', 10, 8, 2),
             ("PR-14-G", "Pressure Regulator with Gauge, 1/4 NPT", 2, 2, 0)]
    rows = "".join(f"<tr><td>{c}</td><td>{e(d)}</td><td class='num'>{o}</td><td class='num'>{s}</td><td class='num'>{b}</td></tr>" for c, d, o, s, b in items)
    body = f"""
<style>{BASE_CSS} body {{ font-size: 12px; }} .items th {{ border-bottom: 2px solid #234; padding: 6px; text-align: left; }} .items td {{ padding: 6px; border-bottom: 1px solid #ccd; }}</style>
<table><tr><td><div style='font-size:20px;font-weight:bold;color:#234'>Bayside Industrial Supply Co.</div>1850 Zanker Road, Suite 140, San Jose, CA 95112</td>
<td style='text-align:right'><div style='font-size:24px;letter-spacing:2px;color:#234'>PACKING SLIP</div></td></tr></table>
<br>
<table><tr><td style='width:50%'><b>SHIP TO</b><br>Northwind Manufacturing Inc.<br>Dock 4, 2200 Mission College Blvd<br>Santa Clara, CA 95054</td>
<td><table><tr><td class='b'>Packing Slip #</td><td class='num'>PS-10442</td></tr><tr><td class='b'>Ship Date</td><td class='num'>09/30/2026</td></tr>
<tr><td class='b'>PO Number</td><td class='num'>PO-7731</td></tr><tr><td class='b'>Carrier</td><td class='num'>UPS Ground</td></tr></table></td></tr></table>
<br>
<table class='items'><tr><th>Item #</th><th>Description</th><th class='num'>Qty Ordered</th><th class='num'>Qty Shipped</th><th class='num'>Backordered</th></tr>{rows}</table>
<p>Backordered items will ship separately. Please check the contents against this slip on arrival.</p>
"""
    t = {"po_number": "PO-7731", "document_number": "PS-10442", "receipt_date": "2026-09-30", "receiver_name": None,
         "line_items": [{"description": d, "quantity": float(s)} for _, d, _, s, _ in items]}
    return Layout("gr_packing_slip", "US packing slip (ordered / shipped / backordered)", body, t, forms=("digital.pdf", "photo.jpg"))


def build_goods_receipts(browser: str, rng) -> None:
    build_set(browser, rng, GR_OUT, (gr_challan, gr_grn, gr_packing_slip))


def build_purchase_orders(browser: str, rng) -> None:
    build_set(browser, rng, PO_OUT, (po_indian, po_us))


def build_set(browser: str, rng, out: Path, makers) -> None:
    for make in makers:
        lay = make()
        folder = out / lay.name
        if folder.exists():
            shutil.rmtree(folder)
        folder.mkdir(parents=True)
        digital = folder / "digital.pdf"
        print_pdf(browser, lay.html, digital)
        for form in lay.forms:
            if form != "digital.pdf":
                FORM_WRITERS[form](digital, folder / form, rng)
        (folder / "truth.json").write_text(json.dumps(lay.truth, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        print(f"{lay.name}: {', '.join(lay.forms)}")


if __name__ == "__main__":
    main()
